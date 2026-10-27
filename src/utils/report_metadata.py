"""Traceable reporting of recorded run values, without inventing missing data."""
from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path

import pandas as pd


def file_hash(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def attach_report_evidence(row, payload, metrics_path):
    """Carry the validated payload and bind subsequent reporting to its files."""
    path = Path(metrics_path).resolve()
    if json.loads(path.read_text()) != payload:
        raise RuntimeError(f"Metrics changed during validation: {path}")
    count = payload.get("metrics", {}).get("num_params")
    if count is not None and (isinstance(count, bool) or not isinstance(count, (int, float))
                              or not math.isfinite(count) or count < 0 or int(count) != count):
        raise RuntimeError(f"Invalid recorded parameter count: {path}")
    measured = payload.get("model_measurements") or {}
    if "tensors" in measured:
        tensors = measured["tensors"]
        total = sum(t["elements"] for t in tensors)
        trainable = sum(t["elements"] for t in tensors if t["trainable"])
        if (len({t["name"] for t in tensors}) != len(tensors)
                or any(math.prod(t["shape"]) != t["elements"] for t in tensors)
                or measured.get("total_parameters") != total
                or measured.get("trainable_parameters") != trainable
                or measured.get("frozen_parameters") != total - trainable
                or (count is not None and count not in {total, trainable})):
            raise RuntimeError(f"Inconsistent measured parameter inventory: {path}")
    files = [path, path.parent / payload["predictions_file"]]
    row["_evidence"] = {
        "path": path, "payload": payload,
        "hashes": {str(p): file_hash(p) for p in files},
    }
    return row


def verify_report_snapshot(rows):
    for row in rows:
        for name, expected in row.get("_evidence", {}).get("hashes", {}).items():
            if not Path(name).is_file() or file_hash(name) != expected:
                raise RuntimeError(f"Report input changed during build: {name}")


def flatten(value, prefix=""):
    """JSON leaves as typed cells; lists use indexed paths to avoid truncation."""
    if isinstance(value, dict) and value:
        for key, child in value.items():
            yield from flatten(child, f"{prefix}.{key}" if prefix else str(key))
    elif isinstance(value, list) and value:
        for index, child in enumerate(value):
            yield from flatten(child, f"{prefix}[{index}]")
    else:
        if value is None:
            yield prefix, "not recorded", "null"
        elif isinstance(value, (dict, list)):
            yield prefix, json.dumps(value), type(value).__name__
        elif isinstance(value, float) and not math.isfinite(value):
            yield prefix, str(value), "nonfinite recorded value"
        else:
            yield prefix, value, type(value).__name__


def reporting_sheets(experiment_rows):
    """All saved evaluation/config values plus explicitly labelled log summaries.

    Weight tensors stay in checkpoints. Epoch history stays in the linked log:
    extrema and last values are summaries, never substituted for TEST metrics.
    """
    details, values, files, tensors = [], [], [], []
    for experiment, rows in experiment_rows.items():
        for row in rows:
            evidence = row.get("_evidence")
            if evidence is None:  # callers providing synthetic summary rows
                continue
            path, payload = evidence["path"], evidence["payload"]
            identity = dict(experiment=experiment, run=path.parent.name,
                            model=row["model"], variant=row["variant"], seed=row["seed"])
            completion = payload.get("training_completion") or {}
            count = payload.get("metrics", {}).get("num_params")
            contract = payload.get("run_contract") or {}
            model = contract.get("model_config") or {}
            training = contract.get("training_config") or {}
            data = contract.get("data_config") or {}
            network = model.get("model", model)
            layer = network.get("layer") or {}
            optimizer = training.get("optimizer")
            optimizer_options = optimizer if isinstance(optimizer, dict) else training
            measured = payload.get("model_measurements") or {}
            tree = row["model"] in {"xgboost", "lightgbm", "random_forest"}
            baseline = row["model"] in {"random_walk", "seasonal_naive", "moving_average"}
            neural = not (tree or baseline)
            details.append({
                **identity, "metrics_path": str(path),
                "num_params": count,
                "total_parameters": measured.get("total_parameters"),
                "trainable_parameters": measured.get("trainable_parameters"),
                "frozen_parameters": measured.get("frozen_parameters"),
                "parameter_bytes": measured.get("parameter_bytes"),
                "hidden_size": network.get("hidden_size", network.get("d_model")) if neural else None,
                "layers": network.get("num_layers", network.get("n_layers")) if neural else None,
                "state_size": network.get("d_state", layer.get("d_state")) if neural else None,
                "dropout": network.get("dropout") if neural else None,
                "optimizer": (optimizer.get("_name_") if isinstance(optimizer, dict) else optimizer)
                    if neural else ("native tree estimator" if tree else "not applicable"),
                "initial_lr": optimizer_options.get("lr") if neural else
                    (model.get("learning_rate") if row["model"] in {"xgboost", "lightgbm"} else None),
                "weight_decay": optimizer_options.get("weight_decay") if neural else None,
                "n_estimators_configured": model.get("n_estimators") if tree else None,
                "tree_max_depth": model.get("max_depth") if tree else None,
                "batch_size_configured": training.get("batch_size", (training.get("loader_train") or {}).get("batch_size")),
                "input_months": data.get("input_len"), "lag_count": data.get("lag_count"),
                "train_end": data.get("train_end"), "val_start": data.get("val_start"),
                "val_end": data.get("val_end"), "test_start": data.get("test_start"),
                "test_end": data.get("test_end"),
                "parameter_count_status": "recorded" if count is not None else
                    ("not applicable: native tree/baseline" if row["model"] in
                     {"xgboost", "lightgbm", "random_forest", "random_walk",
                      "seasonal_naive", "moving_average"} else "not recorded"),
                "n_test_targets": payload.get("n_samples"),
                "epochs_run": completion.get("epochs_run"),
                "best_epoch_recorded": completion.get("best_epoch"),
                "best_validation_loss": completion.get("best_val"),
                "stop_reason": completion.get("stop_reason"),
                "source_fingerprint": payload.get("source_fingerprint"),
                "checkpoint": payload.get("checkpoint"),
                "checkpoint_sha256": completion.get("checkpoint_sha256"),
                "bidirectional": (payload.get("architecture") or {}).get("bidirectional"),
            })
            # The completion repeats the exact contract; keep one readable copy.
            saved = dict(payload)
            measurements = saved.get("model_measurements") or {}
            for tensor in measurements.get("tensors", []):
                tensors.append({**identity, **tensor,
                                "shape": json.dumps(tensor["shape"])})
            if measurements:
                saved["model_measurements"] = {k: v for k, v in measurements.items() if k != "tensors"}
            if isinstance(saved.get("training_completion"), dict):
                saved["training_completion"] = {k: v for k, v in completion.items()
                                                if k != "run_contract"}
            sources = [(path, saved, "validated evaluation/stored contract; inherited unused settings may remain")]
            for relative in ("logs/status.json", "hydra/logs/status.json",
                             "hydra/logs/resolved_config.json", "baseline_definition.json",
                             "checkpoint_manifest.json", "hydra/checkpoint_manifest.json",
                             "best_xgb_metadata.json"):
                extra = path.parent / relative
                if extra.is_file():
                    raw = extra.read_bytes()
                    evidence["hashes"][str(extra)] = hashlib.sha256(raw).hexdigest()
                    sources.append((extra, json.loads(raw), "recorded auxiliary metadata"))
            # Explicitly identify log scope: an external log can contain multiple
            # attempts; these extrema are not a selected checkpoint's test score.
            for relative in ("logs/metrics.json", "hydra/logs/epochs.jsonl"):
                log = path.parent / relative
                if not log.is_file():
                    continue
                raw = log.read_bytes()
                evidence["hashes"][str(log)] = hashlib.sha256(raw).hexdigest()
                records = ([json.loads(line) for line in raw.decode().splitlines() if line.strip()]
                           if log.suffix == ".jsonl" else json.loads(raw))
                if not isinstance(records, list):
                    raise ValueError(f"Expected epoch record list: {log}")
                summary = {"record_count": len(records), "scope": "all records in saved log"}
                columns = {k for record in records for k in record}
                for key in sorted(columns):
                    series = [r[key] for r in records if isinstance(r.get(key), (int, float))
                              and not isinstance(r[key], bool) and math.isfinite(r[key])]
                    if series:
                        summary[key] = dict(min=min(series), max=max(series), last=series[-1],
                                            mean=sum(series) / len(series), recorded_count=len(series))
                sources.append((log, summary, "derived log summary; not test metrics"))
            for source, data, status in sources:
                for key, value, kind in flatten(data):
                    values.append({**identity, "source": str(source), "evidence": status,
                                   "field": key, "value": value, "type": kind})
            for source, digest in evidence["hashes"].items():
                files.append({**identity, "path": source, "sha256": digest})
    return {"Run_Details": pd.DataFrame(details), "Run_Values": pd.DataFrame(values),
            "Model_Tensors": pd.DataFrame(tensors),
            "Report_Files": pd.DataFrame(files)}


def metric_dictionary():
    from src.inference.metrics import REQUIRED_EVALUATION_METRICS
    rows = []
    for name in REQUIRED_EVALUATION_METRICS:
        if "mape" in name:
            units = "percent, physical USD/kg targets"
            rule = ("MAPE excludes absolute targets <= 1 in each channel's physical units"
                    if "smape" not in name else "sMAPE: 0/0 contributes zero; range 0..200%")
        elif "mase" in name or "rmsse" in name:
            units = "dimensionless, train-only seasonal scale (12 months)"
            rule = "Scale within each series, then mean across series and the two channels"
        else:
            units = "dimensionless, train-fitted per-series min-max normalization"
            rule = "Both channels equally weighted; valid targets only"
            if name == "rmse":
                rule = "Mean of channel RMSEs; not sqrt of combined MSE"
            elif "r2" in name:
                rule = "Per-channel R2 across targets; combined is the channel mean"
        rows.append(dict(metric=name, units=units, definition=rule))
    return pd.DataFrame(rows)


def verify_workbook(path, sheets):
    """Read back every stored cell before publishing; Excel has ~15-digit precision."""
    from openpyxl import load_workbook
    with Path(path).open("rb") as stream:
        book = load_workbook(stream, read_only=True, data_only=False)
        try:
            for name, frame in sheets.items():
                for offset in range(0, max(1, len(frame)), 1_000_000):
                    sheet_name = name if offset == 0 else f"{name}_{offset // 1_000_000 + 1}"
                    chunk = frame.iloc[offset:offset + 1_000_000]
                    ws = book[sheet_name]
                    rows = ws.iter_rows()
                    header = [c.value for c in next(rows)]
                    if frame.empty:
                        continue
                    if header != list(frame.columns):
                        raise RuntimeError(f"Workbook column mismatch: {sheet_name}")
                    count = 0
                    for expected in chunk.itertuples(index=False, name=None):
                        actual = next(rows, None)
                        if actual is None:
                            raise RuntimeError(f"Workbook missing row: {sheet_name}")
                        for cell, wanted in zip(actual, expected):
                            if cell.data_type == "f":
                                raise RuntimeError(f"Unexpected formula: {sheet_name}/{cell.coordinate}")
                            if pd.isna(wanted) or wanted == "":
                                valid = cell.value is None
                            elif isinstance(wanted, (int, float)) and not isinstance(wanted, bool):
                                valid = isinstance(cell.value, (int, float)) and math.isclose(
                                    cell.value, wanted, rel_tol=1e-14, abs_tol=0.0)
                            else:
                                valid = cell.value == wanted
                            if not valid:
                                raise RuntimeError(f"Workbook cell mismatch: {sheet_name}/{cell.coordinate}")
                        count += 1
                    if count != len(chunk) or next(rows, None) is not None:
                        raise RuntimeError(f"Workbook row count mismatch: {sheet_name}")
        finally:
            book.close()
